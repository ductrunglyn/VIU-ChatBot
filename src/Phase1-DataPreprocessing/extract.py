"""Bước 1 & 2: Trích xuất văn bản (có nhận biết bảng) từ PDF / Word / Excel.

Kết quả: mỗi file gốc -> 1 file Markdown thô trong data/interim/, giữ lại
ranh giới trang bằng dấu mốc <<<PAGE n>>> để bước làm sạch nhận diện đầu/chân trang.
"""
from __future__ import annotations
import re
from pathlib import Path
from typing import List

PAGE_MARKER = "<<<PAGE {n}>>>"


def _PAGE_STRIP(text: str) -> str:
    """Bỏ dấu mốc PAGE và khoảng trắng để kiểm tra văn bản có thực sự rỗng không."""
    return re.sub(r"<<<PAGE \d+>>>", "", text).replace("\n", "").strip()


def _table_to_markdown(table: List[List]) -> str:
    """Chuyển 1 bảng (list các hàng) sang bảng Markdown."""
    rows = [[("" if c is None else str(c).replace("\n", " ").strip()) for c in row]
            for row in table if row]
    if not rows:
        return ""
    ncol = max(len(r) for r in rows)
    rows = [r + [""] * (ncol - len(r)) for r in rows]  # đệm cho đủ cột
    header = rows[0]
    body = rows[1:]
    md = ["| " + " | ".join(header) + " |",
          "| " + " | ".join(["---"] * ncol) + " |"]
    for r in body:
        md.append("| " + " | ".join(r) + " |")
    return "\n".join(md)


def extract_pdf(path: Path, use_ocr: bool = True, min_chars: int = 30) -> str:
    """Trích text + bảng từ PDF theo từng trang.

    Trang nào có ít hơn `min_chars` ký tự (tức là trang scan ảnh) sẽ được OCR
    nếu use_ocr=True. Nhờ vậy file 'hỗn hợp' (vừa text vừa scan) vẫn xử lý tối ưu.
    """
    import pdfplumber

    path = Path(path)  # chấp nhận cả chuỗi lẫn Path

    pages_out = []   # mỗi phần tử: [số trang, text, [markdown bảng...]]
    scanned = []     # chỉ số (1-based) các trang scan đã OCR
    fitz_doc = None  # mở lười, chỉ khi cần OCR
    n_ocr = 0
    try:
        with pdfplumber.open(path) as pdf:
            for i, page in enumerate(pdf.pages, start=1):
                text = (page.extract_text() or "").strip()
                # bảng từ pdfplumber (chỉ có với trang text số hóa)
                tbls_md = [md for tbl in page.extract_tables()
                           if (md := _table_to_markdown(tbl))]
                if len(text) >= min_chars:
                    pages_out.append([i, text, tbls_md])
                elif use_ocr:
                    if fitz_doc is None:
                        import fitz
                        import ocr
                        fitz_doc = fitz.open(path)
                        _ocr = ocr
                    ocr_text = _ocr.ocr_page(fitz_doc[i - 1])
                    pages_out.append([i, ocr_text, tbls_md])
                    if ocr_text:
                        n_ocr += 1
                    scanned.append(i)
                else:
                    pages_out.append([i, text, tbls_md])
    finally:
        if fitz_doc is not None:
            fitz_doc.close()

    # OCR BẢNG cho các trang scan (pdfplumber không đọc được bảng ảnh) bằng img2table.
    if use_ocr and scanned:
        try:
            import table_ocr
            ptables = table_ocr.extract_tables_markdown(path)  # {chỉ số 0-based: md}
            n_tbl = 0
            for entry in pages_out:
                idx0 = entry[0] - 1
                if entry[0] in scanned and idx0 in ptables:
                    entry[2].append(ptables[idx0])
                    n_tbl += 1
            if n_tbl:
                print(f"    [TABLE] Trích {n_tbl} trang có bảng (scan) trong '{path.name}'.")
        except Exception as e:  # noqa: BLE001 - không để lỗi bảng chặn cả file
            print(f"    ⚠️  Bỏ qua OCR bảng cho '{path.name}': {e}")

    out: List[str] = []
    for i, text, tbls_md in pages_out:
        out.append(PAGE_MARKER.format(n=i))
        if text:
            out.append(text)
        for md in tbls_md:
            out.append("\n" + md + "\n")

    if n_ocr:
        print(f"    [OCR] Đã OCR {n_ocr} trang scan trong '{path.name}'.")
    body = "\n".join(out).strip()
    if len(_PAGE_STRIP(body)) < 20:
        print(f"    ⚠️  '{path.name}' vẫn rỗng sau xử lý -> kiểm tra lại file.")
    return body


def _iter_docx_blocks(doc):
    """Duyệt đoạn văn và bảng THEO ĐÚNG THỨ TỰ xuất hiện trong file.

    python-docx cho sẵn doc.paragraphs và doc.tables nhưng là hai danh sách rời,
    mất thứ tự xen kẽ. Với kế hoạch đào tạo thì đó là lỗi chí mạng: tài liệu viết
    "Học kỳ 1:" rồi tới bảng học phần của kỳ 1, "Học kỳ 2:" rồi bảng kỳ 2... Xuất
    hết tiêu đề trước rồi mới xuất 11 bảng liền nhau thì không còn biết bảng nào
    thuộc kỳ nào.
    """
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    for child in doc.element.body.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, doc)
        elif isinstance(child, CT_Tbl):
            yield Table(child, doc)


def _dedupe_merged(row: List[str]) -> List[str]:
    """Ô gộp bị python-docx lặp lại text ra mọi cột ("Bắt buộc | Bắt buộc | 15").

    Giữ lần xuất hiện đầu, các ô lặp liền kề để trống cho bảng dễ đọc.
    """
    out = []
    for i, cell in enumerate(row):
        out.append("" if i > 0 and cell == row[i - 1] else cell)
    return out


def extract_docx(path: Path) -> str:
    """Trích text + bảng từ file Word, giữ nguyên thứ tự tiêu đề - bảng."""
    from docx import Document
    from docx.table import Table

    doc = Document(path)
    out: List[str] = []
    for block in _iter_docx_blocks(doc):
        if isinstance(block, Table):
            rows = [_dedupe_merged([c.text.strip() for c in r.cells]) for r in block.rows]
            md = _table_to_markdown(rows)
            if md:
                out.append("\n" + md + "\n")
        elif block.text.strip():
            out.append(block.text.strip())
    return "\n".join(out).strip()


def extract_excel(path: Path) -> str:
    """Mỗi sheet Excel -> 1 bảng Markdown, kèm tên sheet làm tiêu đề."""
    import pandas as pd

    out: List[str] = []
    xls = pd.read_excel(path, sheet_name=None, header=None, dtype=str)
    for sheet, df in xls.items():
        df = df.fillna("")
        out.append(f"## Sheet: {sheet}")
        rows = df.values.tolist()
        md = _table_to_markdown(rows)
        if md:
            out.append(md)
    return "\n\n".join(out).strip()


def extract_doc(path: Path) -> str:
    """Trích văn bản từ tệp Word 97-2003 (.doc, định dạng nhị phân OLE2)."""
    from doc_legacy import extract_doc_text
    return extract_doc_text(path)


def extract_file(path: Path) -> str:
    ext = path.suffix.lower()
    if ext == ".pdf":
        return extract_pdf(path)
    if ext == ".docx":
        return extract_docx(path)
    if ext == ".doc":
        return extract_doc(path)
    if ext in {".xlsx", ".xls"}:
        return extract_excel(path)
    if ext in {".txt", ".md"}:
        return path.read_text(encoding="utf-8", errors="ignore").strip()
    raise ValueError(f"Định dạng chưa hỗ trợ: {ext}")
