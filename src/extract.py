"""Bước 1 & 2: Trích xuất văn bản (có nhận biết bảng) từ PDF / Word / Excel.

Kết quả: mỗi file gốc -> 1 file Markdown thô trong data/interim/, giữ lại
ranh giới trang bằng dấu mốc <<<PAGE n>>> để bước làm sạch nhận diện đầu/chân trang.
"""
from __future__ import annotations
import re
from pathlib import Path
from typing import List

PAGE_MARKER = "<<<PAGE {n}>>>"


def _PAGE_STRIP(text: str) -> str:
    """Bỏ dấu mốc PAGE và khoảng trắng để kiểm tra văn bản có thực sự rỗng không."""
    return re.sub(r"<<<PAGE \d+>>>", "", text).replace("\n", "").strip()


def _table_to_markdown(table: List[List]) -> str:
    """Chuyển 1 bảng (list các hàng) sang bảng Markdown."""
    rows = [[("" if c is None else str(c).replace("\n", " ").strip()) for c in row]
            for row in table if row]
    if not rows:
        return ""
    ncol = max(len(r) for r in rows)
    rows = [r + [""] * (ncol - len(r)) for r in rows]  # đệm cho đủ cột
    header = rows[0]
    body = rows[1:]
    md = ["| " + " | ".join(header) + " |",
          "| " + " | ".join(["---"] * ncol) + " |"]
    for r in body:
        md.append("| " + " | ".join(r) + " |")
    return "\n".join(md)


def extract_pdf(path: Path, use_ocr: bool = True, min_chars: int = 30) -> str:
    """Trích text + bảng từ PDF theo từng trang.

    Trang nào có ít hơn `min_chars` ký tự (tức là trang scan ảnh) sẽ được OCR
    nếu use_ocr=True. Nhờ vậy file 'hỗn hợp' (vừa text vừa scan) vẫn xử lý tối ưu.
    """
    import pdfplumber

    out: List[str] = []
    fitz_doc = None  # mở lười, chỉ khi cần OCR
    n_ocr = 0
    try:
        with pdfplumber.open(path) as pdf:
            for i, page in enumerate(pdf.pages, start=1):
                out.append(PAGE_MARKER.format(n=i))
                text = (page.extract_text() or "").strip()
                if len(text) >= min_chars:
                    out.append(text)
                elif use_ocr:
                    if fitz_doc is None:
                        import fitz
                        import ocr
                        fitz_doc = fitz.open(path)
                        _ocr = ocr
                    ocr_text = _ocr.ocr_page(fitz_doc[i - 1])
                    if ocr_text:
                        out.append(ocr_text)
                        n_ocr += 1
                elif text:
                    out.append(text)
                # Bảng: quan trọng với khung chương trình đào tạo
                for tbl in page.extract_tables():
                    md = _table_to_markdown(tbl)
                    if md:
                        out.append("\n" + md + "\n")
    finally:
        if fitz_doc is not None:
            fitz_doc.close()

    if n_ocr:
        print(f"    [OCR] Đã OCR {n_ocr} trang scan trong '{path.name}'.")
    body = "\n".join(out).strip()
    if len(_PAGE_STRIP(body)) < 20:
        print(f"    ⚠️  '{path.name}' vẫn rỗng sau xử lý -> kiểm tra lại file.")
    return body


def extract_docx(path: Path) -> str:
    """Trích text + bảng từ file Word."""
    from docx import Document

    doc = Document(path)
    out: List[str] = []
    for para in doc.paragraphs:
        if para.text.strip():
            out.append(para.text)
    for tbl in doc.tables:
        rows = [[cell.text for cell in row.cells] for row in tbl.rows]
        md = _table_to_markdown(rows)
        if md:
            out.append("\n" + md + "\n")
    return "\n".join(out).strip()


def extract_excel(path: Path) -> str:
    """Mỗi sheet Excel -> 1 bảng Markdown, kèm tên sheet làm tiêu đề."""
    import pandas as pd

    out: List[str] = []
    xls = pd.read_excel(path, sheet_name=None, header=None, dtype=str)
    for sheet, df in xls.items():
        df = df.fillna("")
        out.append(f"## Sheet: {sheet}")
        rows = df.values.tolist()
        md = _table_to_markdown(rows)
        if md:
            out.append(md)
    return "\n\n".join(out).strip()


def extract_file(path: Path) -> str:
    ext = path.suffix.lower()
    if ext == ".pdf":
        return extract_pdf(path)
    if ext == ".docx":
        return extract_docx(path)
    if ext in {".xlsx", ".xls"}:
        return extract_excel(path)
    if ext in {".txt", ".md"}:
        return path.read_text(encoding="utf-8", errors="ignore").strip()
    raise ValueError(f"Định dạng chưa hỗ trợ: {ext}")
